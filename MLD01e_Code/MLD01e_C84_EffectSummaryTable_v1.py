"""Build five independent effect-summary tables for Reviewer 2 Comment 14.

The output separates the overall sample from the four subgroup analyses:
Overall sample, Gender, Age, Income, and Location. Each table reports only
model-estimated total, direct, and indirect effects. Formal focal-minus-
reference inference remains in the separate subgroup-difference table.
Spatial heterogeneity is intentionally excluded and remains represented by
Figure 8.
"""

from __future__ import annotations

from pathlib import Path
from shutil import copy2
from textwrap import fill

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

import LoadDataFromDisk

matplotlib.use("Agg")
import matplotlib.pyplot as plt


PROJECT_ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = PROJECT_ROOT / "MLD01e_Results" / "MLD01e1_Results"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "comment14-effect-summary"

COMBINED_NAME = "Table_EffectSummary"
COMBINED_CANONICAL_XLSX = RESULTS_DIR / f"{COMBINED_NAME}.xlsx"
COMBINED_CANONICAL_CSV = RESULTS_DIR / f"{COMBINED_NAME}.csv"
COMBINED_DELIVERY_XLSX = OUTPUT_DIR / f"{COMBINED_NAME}.xlsx"
COMBINED_DELIVERY_CSV = OUTPUT_DIR / f"{COMBINED_NAME}.csv"
OVERVIEW_PNG = OUTPUT_DIR / f"{COMBINED_NAME}_preview.png"

SUBGROUP_RESULTS = RESULTS_DIR / "Table_SubgroupDifferenceInference.csv"
EXPECTED_N = 11_568
INCOME_THRESHOLD = 170_000.0

OUTCOMES = {
    "Climate knowledge": (
        "factual_knowledge_probs.parquet",
        "counterfactual_knowledge_probs.parquet",
    ),
    "Climate awareness": (
        "factual_awareness_probs.parquet",
        "counterfactual_awareness_probs.parquet",
    ),
    "Soil and water conservation": (
        "factual_action_probs_soilandwater.parquet",
        "counterfactual_action_probs_soilandwater.parquet",
    ),
    "Risk reduction": (
        "factual_action_probs_riskreduction.parquet",
        "counterfactual_action_probs_riskreduction.parquet",
    ),
    "Road improvement": (
        "factual_action_probs_roadimprovement.parquet",
        "counterfactual_action_probs_roadimprovement.parquet",
    ),
    "Community participation": (
        "factual_action_probs_communityparticipation.parquet",
        "counterfactual_action_probs_communityparticipation.parquet",
    ),
}

TABLE_SPECS = {
    "Overall sample": {
        "slug": "OverallSample",
        "title": "Summary of model-estimated total, direct, and indirect effects in the overall sample",
    },
    "Gender": {
        "slug": "Gender",
        "title": "Summary of model-estimated total, direct, and indirect effects by gender",
        "focal": "Female",
        "reference": "Male",
        "contrast": "Female − Male",
        "indicator": "female",
    },
    "Age": {
        "slug": "Age",
        "title": "Summary of model-estimated total, direct, and indirect effects by age",
        "focal": "Age >65",
        "reference": "Age 45–65",
        "contrast": "Age >65 − Age 45–65",
        "indicator": "older_than_65",
    },
    "Income": {
        "slug": "Income",
        "title": "Summary of model-estimated total, direct, and indirect effects by income",
        "focal": "Low income",
        "reference": "High income",
        "contrast": "Low income (≤170,000) − High income (>170,000)",
        "indicator": "low_income",
    },
    "Location": {
        "slug": "Location",
        "title": "Summary of model-estimated total, direct, and indirect effects by location",
        "focal": "Rural",
        "reference": "Urban",
        "contrast": "Rural − Urban",
        "indicator": "rural",
    },
}

COMMON_NOTE = (
    "Note: All values are percentage points. Direct and indirect effects are not "
    "applicable to climate knowledge because it is the first analytical layer. Total, "
    "direct, and indirect contrasts are estimated under separate nonlinear prediction "
    "scenarios and are not constrained to sum exactly."
)

EXPECTED_OVERALL_EFFECTS = {
    "Climate knowledge": (4.720115, None, None),
    "Climate awareness": (1.167906, 0.375852, 0.808328),
    "Soil and water conservation": (1.148910, 0.431733, 0.658283),
    "Risk reduction": (0.563882, 0.022734, 0.537009),
    "Road improvement": (0.747151, -0.187590, 0.915003),
    "Community participation": (0.877488, -0.058883, 0.958994),
}


def _read_probability(path: Path) -> pd.DataFrame:
    frame = pd.read_parquet(path)
    expected_columns = [str(index) for index in range(10)] + ["mean"]
    if list(frame.columns) != expected_columns:
        raise ValueError(f"Unexpected probability columns in {path}: {list(frame.columns)}")
    if len(frame) != EXPECTED_N:
        raise ValueError(f"Unexpected row count in {path}: {len(frame)}")
    repeated_mean = frame[[str(index) for index in range(10)]].mean(axis=1)
    if not np.allclose(frame["mean"], repeated_mean, rtol=0.0, atol=1e-12):
        raise ValueError(f"Stored mean does not equal ten repeated predictions: {path}")
    return frame


def _load_effects() -> dict[str, dict[str, pd.Series]]:
    effects: dict[str, dict[str, pd.Series]] = {}
    for outcome, (factual_name, counterfactual_name) in OUTCOMES.items():
        effects[outcome] = {}
        for scenario in ("total", "direct", "indirect"):
            directory = RESULTS_DIR if scenario == "total" else RESULTS_DIR / scenario
            factual = _read_probability(directory / factual_name)
            counterfactual = _read_probability(directory / counterfactual_name)
            effect = (counterfactual["mean"] - factual["mean"]) * 100
            if effect.isna().any() or not np.isfinite(effect).all():
                raise ValueError(f"Non-finite {scenario} effects for {outcome}")
            effects[outcome][scenario] = effect
    return effects


def _load_groups() -> pd.DataFrame:
    data = LoadDataFromDisk.load_data_from_disk().reset_index()
    if len(data) != EXPECTED_N:
        raise ValueError(f"Unexpected analysis-data row count: {len(data)}")
    groups = pd.DataFrame(
        {
            "female": data["Respon_Female"].astype(int),
            "older_than_65": (data["Respon_Age"] > 65).astype(int),
            "low_income": (data["TotalIncome"] <= INCOME_THRESHOLD).astype(int),
            "rural": data["Rural_Dummy"].astype(int),
        }
    )
    if not all(set(groups[column].unique()).issubset({0, 1}) for column in groups):
        raise ValueError("Subgroup indicators must be binary")
    return groups


def build_tables() -> tuple[dict[str, pd.DataFrame], pd.DataFrame]:
    effects = _load_effects()
    groups = _load_groups()
    inference = pd.read_csv(SUBGROUP_RESULTS)
    if len(inference) != (len(TABLE_SPECS) - 1) * len(OUTCOMES):
        raise ValueError(f"Expected 24 subgroup comparisons, found {len(inference)}")

    tables: dict[str, pd.DataFrame] = {}
    source_rows: list[dict[str, object]] = []

    overall_rows: list[dict[str, object]] = []
    for outcome in OUTCOMES:
        values = {
            scenario: float(effects[outcome][scenario].mean())
            for scenario in ("total", "direct", "indirect")
        }
        if outcome == "Climate knowledge":
            values["direct"] = None
            values["indirect"] = None
        overall_rows.append(
            {
                "Outcome": outcome,
                "Total": values["total"],
                "Direct": values["direct"],
                "Indirect": values["indirect"],
            }
        )
        source_rows.append(
            {
                "Table": "Overall sample",
                "Outcome": outcome,
                "Group": "Overall",
                "Total": values["total"],
                "Direct": values["direct"],
                "Indirect": values["indirect"],
            }
        )
    tables["Overall sample"] = pd.DataFrame(overall_rows)

    for table_name, specification in list(TABLE_SPECS.items())[1:]:
        group = groups[str(specification["indicator"])].astype(bool)
        rows: list[dict[str, object]] = []
        for outcome in OUTCOMES:
            record: dict[str, object] = {"Outcome": outcome}
            for label, mask in (
                (str(specification["focal"]), group),
                (str(specification["reference"]), ~group),
            ):
                values: dict[str, float | None] = {}
                for scenario in ("total", "direct", "indirect"):
                    values[scenario] = (
                        None
                        if outcome == "Climate knowledge" and scenario != "total"
                        else float(effects[outcome][scenario][mask].mean())
                    )
                    record[f"{label} {scenario}"] = values[scenario]
                source_rows.append(
                    {
                        "Table": table_name,
                        "Outcome": outcome,
                        "Group": label,
                        "Total": values["total"],
                        "Direct": values["direct"],
                        "Indirect": values["indirect"],
                    }
                )

            infer = inference.loc[
                inference["outcome"].eq(outcome)
                & inference["contrast"].eq(specification["contrast"])
            ]
            if len(infer) != 1:
                raise ValueError(f"Missing inference row: {table_name}, {outcome}")
            infer_row = infer.iloc[0]
            focal_total = float(record[f"{specification['focal']} total"])
            reference_total = float(record[f"{specification['reference']} total"])
            if not np.isclose(
                focal_total,
                float(infer_row["focal_mean_total_effect_pp"]),
                rtol=0.0,
                atol=1e-12,
            ):
                raise ValueError(f"Focal total does not match inference table: {table_name}, {outcome}")
            if not np.isclose(
                reference_total,
                float(infer_row["reference_mean_total_effect_pp"]),
                rtol=0.0,
                atol=1e-12,
            ):
                raise ValueError(f"Reference total does not match inference table: {table_name}, {outcome}")
            if not np.isclose(
                focal_total - reference_total,
                float(infer_row["difference_pp"]),
                rtol=0.0,
                atol=1e-12,
            ):
                raise ValueError(f"Difference does not match inference table: {table_name}, {outcome}")
            rows.append(record)
        tables[table_name] = pd.DataFrame(rows)

    return tables, pd.DataFrame(source_rows)


def _table_note(table_name: str) -> str:
    if table_name == "Age":
        return COMMON_NOTE + " Age groups are >65 years and 45–65 years."
    if table_name == "Income":
        return COMMON_NOTE + " Low income is ≤170,000, and high income is >170,000."
    return COMMON_NOTE


def _write_sheet(sheet, table_name: str, frame: pd.DataFrame) -> None:
    specification = TABLE_SPECS[table_name]
    thin_black = Side(style="thin", color="000000")
    regular_font = Font(name="Times New Roman", size=10, color="000000")
    bold_font = Font(name="Times New Roman", size=10, bold=True, color="000000")
    last_column = len(frame.columns)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title = sheet.cell(1, 1, str(specification["title"]))
    title.font = bold_font
    title.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for column_index, header in enumerate(frame.columns, start=1):
        cell = sheet.cell(2, column_index, header)
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin_black, bottom=thin_black)

    for row_index, row in enumerate(frame.itertuples(index=False, name=None), start=3):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column_index, "—" if pd.isna(value) else value)
            cell.font = regular_font
            cell.alignment = Alignment(
                horizontal="left" if column_index == 1 else "center",
                vertical="center",
            )
            if column_index > 1 and not pd.isna(value):
                cell.number_format = "0.000"
    for column_index in range(1, last_column + 1):
        sheet.cell(2 + len(frame), column_index).border = Border(bottom=thin_black)

    note_row = 4 + len(frame)
    sheet.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=last_column,
    )
    note = sheet.cell(note_row, 1, _table_note(table_name))
    note.font = regular_font
    note.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    sheet.column_dimensions["A"].width = 31
    for column_index in range(2, last_column + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 16
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 34
    sheet.row_dimensions[note_row].height = 64
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _write_workbook(path: Path, tables: dict[str, pd.DataFrame]) -> None:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    for table_name, frame in tables.items():
        sheet = workbook.create_sheet(table_name)
        _write_sheet(sheet, table_name, frame)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_source_csv(
    tables: dict[str, pd.DataFrame], source_values: pd.DataFrame
) -> None:
    source_values.to_csv(COMBINED_CANONICAL_CSV, index=False, float_format="%.9f")
    for table_name, frame in tables.items():
        specification = TABLE_SPECS[table_name]
        frame.to_csv(
            RESULTS_DIR / f"Table_EffectSummary_{specification['slug']}.csv",
            index=False,
            float_format="%.9f",
        )


def _display_frame(frame: pd.DataFrame) -> pd.DataFrame:
    display = frame.copy()
    for column in display.columns[1:]:
        display[column] = display[column].map(
            lambda value: "—" if pd.isna(value) else f"{float(value):.3f}"
        )
    return display


def _draw_table(axis: plt.Axes, table_name: str, frame: pd.DataFrame) -> None:
    display = _display_frame(frame)
    axis.axis("off")
    labels = ["Outcome"] + [column.replace(" ", "\n", 1) for column in display.columns[1:]]
    widths = [0.25] + [(0.75 / (len(display.columns) - 1))] * (len(display.columns) - 1)
    table = axis.table(
        cellText=display.values,
        colLabels=labels,
        cellLoc="center",
        colLoc="center",
        colWidths=widths,
        bbox=[0, 0.18, 1, 0.68],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(8.5)
    for (row, column), cell in table.get_celld().items():
        cell.set_edgecolor("black")
        cell.set_linewidth(0.8 if row in (0, len(display)) else 0.25)
        cell.set_facecolor("white")
        if row == 0:
            cell.set_text_props(weight="bold")
        elif column == 0:
            cell.set_text_props(ha="left")
    axis.set_title(str(TABLE_SPECS[table_name]["title"]), fontsize=10, fontweight="bold", pad=5)
    axis.text(0, 0.11, fill(_table_note(table_name), width=190), fontsize=7.2, ha="left", va="top")


def _write_previews(tables: dict[str, pd.DataFrame]) -> None:
    preview_paths: list[Path] = []
    for table_name, frame in tables.items():
        specification = TABLE_SPECS[table_name]
        path = OUTPUT_DIR / f"Table_EffectSummary_{specification['slug']}_preview.png"
        figure, axis = plt.subplots(figsize=(13.5, 4.2))
        _draw_table(axis, table_name, frame)
        figure.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
        plt.close(figure)
        preview_paths.append(path)

    figure, axes = plt.subplots(5, 1, figsize=(15, 20))
    for axis, (table_name, frame) in zip(axes, tables.items(), strict=True):
        _draw_table(axis, table_name, frame)
    figure.suptitle("Review overview: five independent effect-summary tables", fontsize=12, fontweight="bold")
    figure.tight_layout(rect=(0.01, 0.01, 0.99, 0.98))
    figure.savefig(OVERVIEW_PNG, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def _copy_delivery_files() -> None:
    copy2(COMBINED_CANONICAL_XLSX, COMBINED_DELIVERY_XLSX)
    copy2(COMBINED_CANONICAL_CSV, COMBINED_DELIVERY_CSV)
    for specification in TABLE_SPECS.values():
        slug = str(specification["slug"])
        for suffix in ("xlsx", "csv"):
            copy2(
                RESULTS_DIR / f"Table_EffectSummary_{slug}.{suffix}",
                OUTPUT_DIR / f"Table_EffectSummary_{slug}.{suffix}",
            )


def _validate_outputs(tables: dict[str, pd.DataFrame], source_values: pd.DataFrame) -> None:
    if list(tables) != list(TABLE_SPECS):
        raise ValueError("Table order changed")
    if tables["Overall sample"].shape != (6, 4):
        raise ValueError(f"Unexpected overall-table shape: {tables['Overall sample'].shape}")
    for table_name in list(TABLE_SPECS)[1:]:
        if tables[table_name].shape != (6, 7):
            raise ValueError(f"Unexpected {table_name} table shape: {tables[table_name].shape}")
    if len(source_values) != 54:
        raise ValueError(f"Unexpected source-value row count: {len(source_values)}")
    if any("difference" in column.lower() for frame in tables.values() for column in frame.columns):
        raise ValueError("Effect-summary tables must not contain a difference column")

    overall = tables["Overall sample"]
    for outcome, expected in EXPECTED_OVERALL_EFFECTS.items():
        record = overall.loc[overall["Outcome"].eq(outcome)].iloc[0]
        for column, expected_value in zip(("Total", "Direct", "Indirect"), expected, strict=True):
            if expected_value is None:
                if not pd.isna(record[column]):
                    raise ValueError(f"Knowledge {column.lower()} must be not applicable")
                continue
            if not np.isclose(float(record[column]), expected_value, rtol=0.0, atol=5e-7):
                raise ValueError(f"Unexpected overall {column.lower()} effect for {outcome}")

    workbook = load_workbook(COMBINED_CANONICAL_XLSX, data_only=False)
    if workbook.sheetnames != list(TABLE_SPECS):
        raise ValueError(f"Unexpected combined-workbook sheets: {workbook.sheetnames}")
    for table_name, frame in tables.items():
        sheet = workbook[table_name]
        expected_columns = len(frame.columns)
        if sheet.max_column != expected_columns or sheet.max_row != 10:
            raise ValueError(
                f"Unexpected workbook dimensions for {table_name}: "
                f"{sheet.max_row}x{sheet.max_column}"
            )
        formulas = [
            cell.value
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        if formulas:
            raise ValueError(f"Unexpected formulas in {table_name}: {formulas}")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    tables, source_values = build_tables()

    _write_workbook(COMBINED_CANONICAL_XLSX, tables)
    for table_name, frame in tables.items():
        specification = TABLE_SPECS[table_name]
        _write_workbook(
            RESULTS_DIR / f"Table_EffectSummary_{specification['slug']}.xlsx",
            {table_name: frame},
        )
    _write_source_csv(tables, source_values)
    _copy_delivery_files()
    _write_previews(tables)
    _validate_outputs(tables, source_values)

    for table_name, frame in tables.items():
        print(f"\n{TABLE_SPECS[table_name]['title']}")
        print(_display_frame(frame).to_string(index=False))
    print(f"\nCombined workbook: {COMBINED_DELIVERY_XLSX}")
    print(f"Combined source CSV: {COMBINED_DELIVERY_CSV}")
    print(f"Review overview: {OVERVIEW_PNG}")


if __name__ == "__main__":
    main()
